# -*- coding: utf-8 -*-
"""Generate room-algorithm comparison slides (old greedy vs solver)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

DARK = RGBColor(0x1F, 0x38, 0x64)
RED = RGBColor(0xC0, 0x39, 0x2B)
GREEN = RGBColor(0x1E, 0x84, 0x49)
GRAY = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_RED = RGBColor(0xFD, 0xED, 0xEC)
LIGHT_GREEN = RGBColor(0xE9, 0xF7, 0xEF)
LIGHT_BLUE = RGBColor(0xEA, 0xF2, 0xF8)
FONT = "Microsoft YaHei"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def style_run(run, size=18, bold=False, color=DARK, italic=False):
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", FONT)


def add_text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """lines: list of (text, size, bold, color) or list of runs [(text,size,bold,color),...]"""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(6)
        runs = line if isinstance(line, list) else [line]
        for text, size, bold, color in runs:
            r = p.add_run()
            r.text = text
            style_run(r, size, bold, color)
    return box


def add_box(slide, x, y, w, h, fill, line_color=None):
    sh = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line_color is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line_color
        sh.line.width = Pt(1.25)
    sh.shadow.inherit = False
    return sh


def box_text(sh, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(0.22)
    tf.margin_right = Inches(0.22)
    tf.margin_top = Inches(0.15)
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(5)
        runs = line if isinstance(line, list) else [line]
        for text, size, bold, color in runs:
            r = p.add_run()
            r.text = text
            style_run(r, size, bold, color)


def title_bar(slide, text, color=DARK):
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(1.0)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    bar.shadow.inherit = False
    tf = bar.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.5)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    style_run(r, 28, True, WHITE)


def new_slide():
    return prs.slides.add_slide(BLANK)


def style_cell(cell, text, size=15, bold=False, color=DARK, fill=None,
               align=PP_ALIGN.LEFT):
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    tf = cell.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    style_run(r, size, bold, color)


# ---------- Slide 1: title ----------
s = new_slide()
bg = slide_bg = s.shapes.add_shape(
    MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height
)
bg.fill.solid()
bg.fill.fore_color.rgb = DARK
bg.line.fill.background()
bg.shadow.inherit = False
add_text(
    s, 1.0, 2.4, 11.3, 1.2,
    [("分房算法：问题与改进方案", 44, True, WHITE)],
    align=PP_ALIGN.CENTER,
)
add_text(
    s, 1.0, 3.7, 11.3, 0.8,
    [("现状：贪心 + 补丁    vs    改进：整体求解器", 24, False, RGBColor(0xBD, 0xD7, 0xEE))],
    align=PP_ALIGN.CENTER,
)
add_text(
    s, 1.0, 6.3, 11.3, 0.6,
    [("MoM Spa 房间管理系统 · 2026 年 8 月", 14, False, RGBColor(0x8E, 0xA9, 0xC9))],
    align=PP_ALIGN.CENTER,
)

# ---------- Slide 2: current algorithm ----------
s = new_slide()
title_bar(s, "现在的算法：贪心（Greedy）+ 补丁")
steps = [
    ("① 排序", "预约按开始时间\n一张一张处理"),
    ("② 贪心分房", "每张立刻挑\n“当下最好”的房\n不回头"),
    ("③ 重平衡", "情侣没房时\n试着挪 1–2 个单人\n腾出双人房"),
    ("④ 冲突清理", "两两检查重叠\n删掉“受害者”\n的分配"),
    ("⑤ 事后修复", "被删的预约\n再试一次分房"),
]
x = 0.5
for i, (t, body) in enumerate(steps):
    fill = LIGHT_BLUE if i < 2 else LIGHT_RED
    b = add_box(s, x, 1.5, 2.35, 2.3, fill)
    box_text(b, [
        (t, 18, True, DARK if i < 2 else RED),
        (body, 13, False, GRAY),
    ])
    x += 2.6
add_text(s, 0.6, 4.3, 12.1, 0.5, [[
    ("红色 3 步全是“出错后的补救”", 16, True, RED),
    ("——补丁只覆盖作者见过的场景", 16, False, GRAY),
]])
b = add_box(s, 0.6, 5.0, 12.1, 1.7, LIGHT_BLUE)
box_text(b, [
    ("关键特征", 16, True, DARK),
    ("• 只看眼前这一单，不考虑后面还有谁要来", 15, False, GRAY),
    ("• 分错了靠“挪一挪”补救，补救能力有限", 15, False, GRAY),
    ("• 代码位置：app/room_assigner.py 的 assign_rooms（约 1300 行）", 13, False, GRAY),
])

# ---------- Slide 3: two key concepts ----------
s = new_slide()
title_bar(s, "两个关键概念")
b = add_box(s, 0.6, 1.5, 5.9, 4.6, LIGHT_BLUE)
box_text(b, [
    ("贪心算法", 22, True, DARK),
    ("", 8, False, GRAY),
    ("想象前台一张一张收预约单：", 15, False, GRAY),
    ("• 每拿到一张，立刻给它挑当下最好的房间", 15, False, GRAY),
    ("• 挑完就不回头", 15, False, GRAY),
    ("", 8, False, GRAY),
    ("问题：2 点把情侣房给了单人客，", 15, False, RED),
    ("3 点情侣来了就没房——眼前的决定堵死后面的路", 15, False, RED),
])
b = add_box(s, 6.85, 1.5, 5.9, 4.6, LIGHT_RED)
box_text(b, [
    ("深度受限", 22, True, RED),
    ("", 8, False, GRAY),
    ("发现情侣没房后，代码会“挪人”补救：", 15, False, GRAY),
    ("• 挪 1 个人腾房 → 能做到 ✓", 15, False, GREEN),
    ("• 连着挪 2 个人 → 能做到 ✓", 15, False, GREEN),
    ("• 需要挪 3 个人 → 代码没写，直接放弃 ✗", 15, True, RED),
    ("", 8, False, GRAY),
    ("“深度”= 最多连着挪几个人。", 15, False, GRAY),
    ("写死为 2，超过 2 步的解法算法“看不见”", 15, False, GRAY),
])
add_text(s, 0.6, 6.4, 12.1, 0.6, [
    ("结果：明明有可行方案，屏幕上还是 UNASSIGNED", 17, True, RED),
], align=PP_ALIGN.CENTER)

# ---------- Slide 4: bug example 1 ----------
s = new_slide()
title_bar(s, "错误案例 1：5 号房明明空着，却显示 UNASSIGNED", RED)
add_text(s, 0.6, 1.2, 12.1, 0.5, [[
    ("起因：Square 传来的时间戳差了 2 秒", 17, True, DARK),
    ("（上一单 4:00:02 结束，下一单 4:00:00 开始）", 15, False, GRAY),
]])
rows = [
    ("①", "贪心分房", "情侣 4:00–5:00 正常分到 5 号房 ✓", GREEN),
    ("②", "冲突检查", "严格比较时间戳 → 误判“重叠了 2 秒” → 报告 5 号房超售", RED),
    ("③", "冲突清理", "删掉情侣的分配 → 情侣变成 UNASSIGNED", RED),
    ("④", "事后修复", "修复条件不满足、没救回来 → 5 号房整段空着，情侣却无房", RED),
]
y = 1.9
for num, stage, desc, c in rows:
    b = add_box(s, 0.6, y, 12.1, 0.95, LIGHT_GREEN if c is GREEN else LIGHT_RED)
    box_text(b, [[
        (f"{num} {stage}    ", 16, True, c),
        (desc, 15, False, GRAY),
    ]], anchor=MSO_ANCHOR.MIDDLE)
    y += 1.08
b = add_box(s, 0.6, y + 0.1, 12.1, 1.0, LIGHT_BLUE)
box_text(b, [[
    ("根源：", 16, True, DARK),
    ("“先分 → 再查 → 再删 → 再救”四道工序各管各的，任何两处对“忙/闲”判断不一致，就会自己推翻自己的正确结果", 15, False, GRAY),
]])

# ---------- Slide 5: bug example 2 ----------
s = new_slide()
title_bar(s, "错误案例 2：有解，但算法“看不见”", RED)
tbl = s.shapes.add_table(5, 3, Inches(0.6), Inches(1.3), Inches(5.9), Inches(2.6)).table
tbl.columns[0].width = Inches(2.2)
tbl.columns[1].width = Inches(2.0)
tbl.columns[2].width = Inches(1.7)
for j, h in enumerate(("预约", "时间", "贪心分到")):
    style_cell(tbl.cell(0, j), h, 14, True, WHITE, DARK)
data = [
    ("单人 S3", "2:00–3:00", "4 号房"),
    ("单人 S2", "2:00–3:15", "2 号房"),
    ("单人 S1", "2:30–3:30", "5 号房"),
    ("情侣 C", "3:00–4:00", "无房 ✗"),
]
for i, row in enumerate(data, start=1):
    for j, v in enumerate(row):
        color = RED if "✗" in v else DARK
        style_cell(tbl.cell(i, j), v, 14, i == 4, color,
                   LIGHT_RED if i == 4 else WHITE)
add_text(s, 0.6, 4.15, 5.9, 1.0, [
    ("6 号房被经理锁定；1、3 号房整下午已占用；", 13, False, GRAY),
    ("02D 需要 0+2 同时空，也不可用", 13, False, GRAY),
])
b = add_box(s, 6.85, 1.3, 5.9, 2.6, LIGHT_RED)
box_text(b, [
    ("补救过程", 17, True, RED),
    ("• 挪 1 步：S1 换房 → 没有空房，失败", 14, False, GRAY),
    ("• 挪 2 步：先挪 S2 给 S1 让位 → S2 也没处去，失败", 14, False, GRAY),
    ("• 挪 3 步：代码没有这个分支 → 直接放弃", 14, True, RED),
])
b = add_box(s, 6.85, 4.15, 5.9, 2.5, LIGHT_GREEN)
box_text(b, [
    ("其实存在的解（需要 3 步）", 17, True, GREEN),
    ("① S3：4 号房 → 0 号房", 14, False, GRAY),
    ("② S2：2 号房 → 4 号房", 14, False, GRAY),
    ("③ S1：5 号房 → 2 号房", 14, False, GRAY),
    ("④ 情侣 C 进 5 号房 ✓", 14, True, GREEN),
])

# ---------- Slide 6: root cause ----------
s = new_slide()
title_bar(s, "问题的共同根源：补丁摞补丁")
items = [
    ("每个补丁只覆盖“见过的事故”",
     "2 秒容差、最多挪 2 步、1 小时 lookahead——这些常量都是针对具体事故调出来的，新场景一来又漏"),
    ("占用状态有 4 份，各管各的",
     "贪心、重平衡、冲突清理、修复各自维护一份“哪个房什么时候忙”，任何细微出入就互相矛盾"),
    ("结果",
     "bug 偶发、反复出现；每修一次代码更厚（assign_rooms 已约 1300 行），下一个补丁更容易踩到上一个"),
]
y = 1.5
for t, body in items:
    b = add_box(s, 0.6, y, 12.1, 1.55, LIGHT_RED)
    box_text(b, [
        (t, 18, True, RED),
        (body, 15, False, GRAY),
    ])
    y += 1.75

# ---------- Slide 7: solver ----------
s = new_slide()
title_bar(s, "改进方案：求解器 —— 一次算出全天方案", GREEN)
b = add_box(s, 0.6, 1.4, 12.1, 1.5, LIGHT_GREEN)
box_text(b, [
    ("核心思路", 17, True, GREEN),
    ("不再一张一张处理。把全天所有预约摊在桌上，像拼拼图一样，一次性整体规划出大家都有房的完整方案。", 15, False, GRAY),
])
b = add_box(s, 0.6, 3.1, 5.9, 3.4, LIGHT_BLUE)
box_text(b, [
    ("规则写成“约束”，写一次就够", 16, True, DARK),
    ("• 同一间房时间不能重叠（含 2 秒容差）", 14, False, GRAY),
    ("• 踩背只能用 1 / 3 / 4 号房", 14, False, GRAY),
    ("• 经理手动指定、已开始的预约不许动", 14, False, GRAY),
    ("• 02D 需要 0 + 2 同时空", 14, False, GRAY),
    ("• 房间优先级、尽量不动上次的分配（软目标）", 14, False, GRAY),
])
b = add_box(s, 6.85, 3.1, 5.9, 3.4, LIGHT_GREEN)
box_text(b, [
    ("为什么能根治 bug", 16, True, GREEN),
    ("• 只要存在可行方案，保证找到", 14, True, GREEN),
    ("  ——不存在“挪几步”的限制", 14, False, GRAY),
    ("• 一道工序出方案，天然无冲突", 14, False, GRAY),
    ("  ——没有“分了再删再救”的流水线", 14, False, GRAY),
    ("• 7 间房、几十单/天 → 毫秒级算完", 14, False, GRAY),
    ("  （回溯搜索或 Google OR-Tools CP-SAT）", 13, False, GRAY),
])

# ---------- Slide 8: comparison table ----------
s = new_slide()
title_bar(s, "新旧对比")
tbl = s.shapes.add_table(6, 3, Inches(0.6), Inches(1.4), Inches(12.1), Inches(4.9)).table
tbl.columns[0].width = Inches(2.6)
tbl.columns[1].width = Inches(4.75)
tbl.columns[2].width = Inches(4.75)
for j, h in enumerate(("", "现状：贪心 + 补丁", "改进：求解器")):
    style_cell(tbl.cell(0, j), h, 16, True, WHITE, DARK, PP_ALIGN.CENTER)
rows = [
    ("决策方式", "一单一单，只看眼前", "全天整体规划"),
    ("出错补救", "手写补丁，最多挪 2 步", "不需要补救——方案天然无冲突"),
    ("“有房却 UNASSIGNED”", "偶发、反复出现", "结构上不可能发生"),
    ("有解保证", "没有（贪心可能漏）", "有解必找到"),
    ("代码", "约 1300 行、4 道工序", "约束写一次，可删掉 3 大段补丁"),
]
for i, (k, old, new) in enumerate(rows, start=1):
    style_cell(tbl.cell(i, 0), k, 14, True, DARK, LIGHT_BLUE)
    style_cell(tbl.cell(i, 1), old, 14, False, RED, LIGHT_RED)
    style_cell(tbl.cell(i, 2), new, 14, False, GREEN, LIGHT_GREEN)

# ---------- Slide 9: rollout ----------
s = new_slide()
title_bar(s, "落地三步（每步都可单独验收）", GREEN)
steps = [
    ("第 1 步：先补测试",
     "用 debugging/ 里真实出错日的数据做 pytest 用例（单人、情侣、02D、锁定、踩背、Room 5 bug 复现日），把现有正确行为“锁住”"),
    ("第 2 步：抽离求解逻辑",
     "从 assign_rooms 里抽出纯算法：输入 = 预约列表 + 锁定集合，输出 = 分配方案；数据库读写留在外面。顺手解决优先级列表在两个文件重复维护的问题"),
    ("第 3 步：替换核心",
     "用求解器替换贪心核心，跑同一套测试对比结果，确认无回归后切换"),
]
y = 1.5
for t, body in steps:
    b = add_box(s, 0.6, y, 12.1, 1.55, LIGHT_GREEN)
    box_text(b, [
        (t, 18, True, GREEN),
        (body, 15, False, GRAY),
    ])
    y += 1.75

out = r"c:\Users\berns\spa\MoM_Room_Jaelyn\doc\room-algorithm-comparison.pptx"
prs.save(out)
print("saved:", out)
